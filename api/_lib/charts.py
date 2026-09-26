"""Embedded Excel charts -> Chart models.

The demo cut reads exactly one thing: the first line chart on the worksheet we
already extract. Other chart types are left alone, not guessed at.

A chart does not hold its own data. It points at cells ("'Sales'!$B$2:$M$2"), so
we read those cells through the Cells extract.py has already built. That way a
chart point says "€154,000" exactly as the grid does, and it carries the ref of
the cell it came from. Only when a reference cannot be followed do we fall back
to the copy of the values Excel caches inside the chart.

Like the rest of extraction, nothing here may raise. A chart we cannot read is a
chart we do not mention; the spreadsheet itself carries on exactly as before.
"""

from collections import Counter
from typing import Any, Callable

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_to_tuple

from .models import Cell, Chart, ChartPoint, ChartSeries

KINDS = {"lineChart": "line"}   # the only chart type the demo cut reads
MAX_CHARTS = 1

# (ref or None, value, display, number format) for one cell a chart points at.
Entry = tuple[str | None, Any, str, str | None]


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def read_charts(formulas, values, sheet_name: str, cells: list[Cell],
                format_value: Callable[[Any, str | None], str]) -> list[Chart]:
    """The supported charts on one worksheet. Never raises; [] when there are none.

    `formulas` and `values` are the two workbooks extract.py loads. `format_value`
    is extract's own number formatter, passed in so the two modules do not import
    each other.
    """
    try:
        sheet = formulas[sheet_name]
        found = getattr(sheet, "_charts", None) or []
    except Exception:
        return []

    by_ref = {cell.ref: cell for cell in cells}
    reader = _Reader(formulas, values, sheet_name, by_ref, format_value)

    charts: list[Chart] = []
    for source in found:
        if len(charts) >= MAX_CHARTS:
            break
        try:
            chart = reader.chart(source, f"chart{len(charts) + 1}")
        except Exception:
            chart = None
        if chart is not None:
            charts.append(chart)
    return charts


class _Reader:
    """Follows a chart's references back into the workbook."""

    def __init__(self, formulas, values, sheet_name, by_ref, format_value):
        self.formulas = formulas
        self.values = values
        self.sheet_name = sheet_name
        self.by_ref = by_ref
        self.format_value = format_value

    # --- the chart ------------------------------------------------------------

    def chart(self, source, chart_id: str) -> Chart | None:
        kind = KINDS.get(getattr(source, "tagname", None))
        if kind is None:
            return None

        series = [s for s in (self.series(s) for s in (source.series or [])) if s]
        # A line with no numbers on it is nothing a sighted reader could see.
        if not any(p.value is not None for s in series for p in s.points):
            return None

        return Chart(
            id=chart_id,
            kind=kind,
            title=self.title_text(getattr(source, "title", None)),
            x_title=self.axis_title(getattr(source, "x_axis", None)),
            y_title=self.axis_title(getattr(source, "y_axis", None)),
            anchor=_anchor(getattr(source, "anchor", None)),
            series=series,
        )

    def series(self, source) -> ChartSeries | None:
        try:
            values_ref = source.val.numRef if source.val is not None else None
        except Exception:
            values_ref = None
        if values_ref is None:
            return None

        entries = self.entries(values_ref)
        if not entries:
            return None
        categories = self.categories(getattr(source, "cat", None))

        # The chart shows every point in one format: the series' usual one, the way
        # Excel's axis follows its source cells. A single oddly formatted cell in
        # the row must not make one point read as "6,900,000.0%".
        formats = Counter(fmt for _, value, _, fmt in entries if _is_number(value) and fmt)
        usual = formats.most_common(1)[0][0] if formats else None

        points = []
        for i, (ref, value, display, _) in enumerate(entries):
            number = float(value) if _is_number(value) else None
            points.append(ChartPoint(
                category=categories[i] if i < len(categories) and categories[i] else str(i + 1),
                value=number,
                display=self.format_value(value, usual) if number is not None else display,
                ref=ref,
            ))

        return ChartSeries(name=self.series_name(getattr(source, "tx", None)),
                           source=getattr(values_ref, "f", None), points=points)

    def categories(self, cat) -> list[str]:
        if cat is None:
            return []
        ref = getattr(cat, "strRef", None) or getattr(cat, "numRef", None)
        if ref is None:
            return []
        return [display for _, _, display, _ in self.entries(ref)]

    def series_name(self, tx) -> str | None:
        if tx is None:
            return None
        if getattr(tx, "strRef", None) is not None:
            names = [d for _, _, d, _ in self.entries(tx.strRef) if d]
            return names[0] if names else None
        return getattr(tx, "v", None) or None

    # --- text ----------------------------------------------------------------

    def title_text(self, title) -> str | None:
        try:
            tx = title.tx if title is not None else None
            if tx is None:
                return None
            if tx.rich is not None:
                lines = ["".join(run.t or "" for run in (p.r or [])) for p in tx.rich.p]
                text = " ".join(line.strip() for line in lines if line.strip())
                return text or None
            if tx.strRef is not None:
                names = [d for _, _, d, _ in self.entries(tx.strRef) if d]
                return " ".join(names) or None
        except Exception:
            return None
        return None

    def axis_title(self, axis) -> str | None:
        return self.title_text(getattr(axis, "title", None)) if axis is not None else None

    # --- following references ---------------------------------------------------

    def entries(self, ref) -> list[Entry]:
        """The cells a chart reference points at, in order. Cache as a fallback."""
        formula = getattr(ref, "f", None)
        if formula:
            try:
                resolved = self.resolve(formula)
                if resolved:
                    return resolved
            except Exception:
                pass
        return self.cached(ref)

    def resolve(self, formula: str) -> list[Entry]:
        sheet_name, (min_col, min_row, max_col, max_row) = range_to_tuple(formula)
        if sheet_name not in self.values.sheetnames:
            return []
        same_sheet = sheet_name == self.sheet_name
        values_sheet, formulas_sheet = self.values[sheet_name], self.formulas[sheet_name]

        out: list[Entry] = []
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                ref = f"{get_column_letter(col)}{row}"
                cell = self.by_ref.get(ref) if same_sheet else None
                if cell is not None:
                    # Exactly what the grid holds for this cell.
                    out.append((ref, cell.value, cell.display, cell.number_format))
                    continue
                raw = values_sheet.cell(row=row, column=col).value
                fmt = formulas_sheet.cell(row=row, column=col).number_format
                if isinstance(raw, str) and raw.startswith("#"):
                    out.append((ref if same_sheet else None, None, raw, fmt))
                else:
                    out.append((ref if same_sheet else None, raw,
                                self.format_value(raw, fmt), fmt))
        return out

    def cached(self, ref) -> list[Entry]:
        """The values Excel stored inside the chart when it last saved the file."""
        cache = getattr(ref, "numCache", None) or getattr(ref, "strCache", None)
        if cache is None or not cache.pt:
            return []
        fmt = getattr(cache, "formatCode", None)
        size = max(getattr(cache, "ptCount", 0) or 0, max(p.idx for p in cache.pt) + 1)
        out: list[Entry] = [(None, None, "", fmt) for _ in range(size)]
        for point in cache.pt:
            value: Any = point.v
            if getattr(ref, "numCache", None) is not None:
                try:
                    value = float(value)
                except (TypeError, ValueError):
                    value = None
            out[point.idx] = (None, value, self.format_value(value, fmt), fmt)
        return out


def _anchor(anchor) -> str | None:
    """"B12:J27" for the cells a chart covers, "B12" when only its corner is known."""
    try:
        if isinstance(anchor, str):
            return anchor
        start = getattr(anchor, "_from", None)
        if start is None:
            return None
        corner = f"{get_column_letter(start.col + 1)}{start.row + 1}"
        end = getattr(anchor, "to", None)
        if end is None:
            return corner
        return f"{corner}:{get_column_letter(end.col + 1)}{end.row + 1}"
    except Exception:
        return None
