"""openpyxl -> SheetModel.

The messy half of the engine. Everything here has to survive a spreadsheet nobody
designed for us: missing headers, notes pasted below the data, theme colours we
cannot resolve, dates, text where numbers belong. Nothing in this module may raise
on a real file -- a crash during the demo is worse than a shrug.

Two passes over the workbook, because openpyxl cannot do both at once:
    data_only=False  ->  formulas, and all the formatting
    data_only=True   ->  the values Excel last calculated, including error strings

openpyxl never calculates anything itself. If a file has never been opened by a
spreadsheet app, its formula cells have no cached value and an intended #DIV/0!
simply will not be there.
"""

import datetime as dt
import re
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .charts import read_charts
from .models import Cell, SheetModel

MAX_ROWS = 400          # keep the payload and the function well inside Vercel's limits
MAX_COLS = 60

MONTHS = {"jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"}

# Words that name what a number measures. The label is only ever a word the author
# actually wrote in the sheet -- we look for these, we never infer one. A sheet that
# never says what its numbers are gets no label, and the app stays quiet about it.
MEASURES = (
    "revenue", "sales", "turnover", "profit", "loss", "margin", "income",
    "cost", "costs", "expenses", "spend", "spending", "budget", "price",
    "salary", "salaries", "wages", "units", "quantity", "volume", "count",
    "headcount", "population", "score", "scores", "rating", "ratings",
    "hours", "minutes", "days", "distance", "weight", "temperature",
    "rainfall", "attendance", "orders", "visits", "clicks", "downloads",
)
WHITE = {"FFFFFFFF", "FFFFFF", "00FFFFFF"}
BLACKISH = {"FF000000", "000000", "FF000000"}


# --- colours -----------------------------------------------------------------

def _rgb(colour) -> str | None:
    """An RGB string, "unknown-non-default", or None. Never raises.

    Theme and indexed colours cannot be resolved without the workbook's theme, so
    they are recorded as present-but-unnamed rather than guessed at or dropped.
    """
    if colour is None:
        return None
    try:
        if getattr(colour, "type", None) == "rgb":
            value = colour.rgb
            return value if isinstance(value, str) else "unknown-non-default"
        if getattr(colour, "type", None) in ("theme", "indexed"):
            return "unknown-non-default"
    except Exception:
        return "unknown-non-default"
    return None


def _fill_of(cell) -> str | None:
    """The author's highlight, if there is one a sighted person could see."""
    try:
        fill = cell.fill
        if fill is None or fill.fill_type != "solid":
            return None
        rgb = _rgb(fill.fgColor)
        # A white fill looks exactly like no fill, so it conveys nothing.
        if rgb is None or rgb.upper() in WHITE or rgb == "00000000":
            return None
        return rgb
    except Exception:
        return None


def _font_colour_of(cell) -> str | None:
    """A recoloured font. Ordinary black text is not a signal."""
    try:
        font = cell.font
        if font is None:
            return None
        rgb = _rgb(font.color)
        if rgb is None or rgb.upper() in BLACKISH:
            return None
        return rgb
    except Exception:
        return None


def _has_border(cell) -> bool:
    try:
        border = cell.border
        return any(getattr(border, side).style
                   for side in ("left", "right", "top", "bottom")
                   if getattr(border, side, None))
    except Exception:
        return False


# --- values ------------------------------------------------------------------

def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _format(value: Any, number_format: str | None) -> str:
    """What the cell looks like on screen, near enough to say out loud.

    Respects the number format where that is easy, and otherwise formats numbers
    sensibly rather than reading "82400.0" aloud.
    """
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%d %b %Y" if isinstance(value, dt.date) else "%d %b %Y %H:%M")
    if not _is_number(value):
        return str(value)

    fmt = (number_format or "").lower()
    try:
        if "%" in fmt:
            places = len(re.search(r"0\.(0+)", fmt).group(1)) if re.search(r"0\.(0+)", fmt) else 1
            return f"{value * 100:,.{places}f}%"
        symbol = next((s for s in ("€", "$", "£", "¥") if s in fmt), "")
        places = len(re.search(r"0\.(0+)", fmt).group(1)) if re.search(r"0\.(0+)", fmt) else None
        if places is None:
            places = 0 if float(value).is_integer() else 2
        return f"{symbol}{value:,.{places}f}"
    except Exception:
        return str(value)


# --- shape -------------------------------------------------------------------

def _find_header_row(grid: list[list[Any]], max_scan: int = 10) -> int:
    """The first non-empty row that is mostly text. 1-based."""
    for r, row in enumerate(grid[:max_scan], start=1):
        filled = [v for v in row if v is not None and str(v).strip() != ""]
        if len(filled) < 2:
            continue
        if sum(isinstance(v, str) for v in filled) >= len(filled) * 0.6:
            return r
    return 1


def _find_label_col(grid: list[list[Any]], header_row: int, max_scan: int = 5) -> int:
    """The first column that is mostly text below the header. 1-based."""
    below = grid[header_row:]
    for c in range(min(max_scan, max((len(r) for r in grid), default=1))):
        filled = [row[c] for row in below if c < len(row)
                  and row[c] is not None and str(row[c]).strip() != ""]
        if len(filled) < 2:
            continue
        if sum(isinstance(v, str) for v in filled) >= len(filled) * 0.6:
            return c + 1
    return 1


def _data_end(grid: list[list[Any]], header_row: int) -> int:
    """Where the table stops: the first blank row after it. 1-based, inclusive.

    Arbitrary sheets often carry notes, legends or a second small table below the
    data. Those are not part of the series and must not be compared against it.
    """
    last = header_row
    for r in range(header_row + 1, len(grid) + 1):
        row = grid[r - 1]
        if any(v is not None and str(v).strip() != "" for v in row):
            last = r
        elif last > header_row:
            break
    return last


def _value_label(*sources: Any) -> str:
    """What the numbers measure, if the sheet says so anywhere.

    Looks through the title and the headers for a word that names a quantity, and
    returns "" when it finds none. Only ever returns a word that is literally
    written in the sheet: "European Sales 2026" gives "sales", and a sheet titled
    "Sheet1" gives nothing at all.
    """
    for source in sources:
        for text in (source if isinstance(source, (list, tuple)) else [source]):
            if not isinstance(text, str):
                continue
            words = re.findall(r"[a-z]+", text.lower())
            for measure in MEASURES:
                if measure in words:
                    return measure
    return ""


def _step_word(headers: list[str]) -> str | None:
    """What one step along a row means, if the headers say so."""
    lowered = [h.strip().lower()[:3] for h in headers if isinstance(h, str)]
    if sum(h in MONTHS for h in lowered) >= 3:
        return "month"
    if sum(bool(re.fullmatch(r"q[1-4]", h)) for h in lowered) >= 3:
        return "quarter"
    return None


# --- the main entry point ----------------------------------------------------

def extract(source, sheet_index: int = 0) -> tuple[SheetModel, str | None]:
    """Read the first worksheet into a SheetModel.

    Returns the model and the step word for the attention engine, which cannot work
    out on its own that a column means a month.
    """
    formulas = load_workbook(source, data_only=False)
    try:
        source.seek(0)
    except Exception:
        pass
    values = load_workbook(source, data_only=True)

    name = formulas.sheetnames[sheet_index]
    sf, sv = formulas[name], values[name]

    n_rows = min(sv.max_row or 1, MAX_ROWS)
    n_cols = min(sv.max_column or 1, MAX_COLS)
    grid = [[sv.cell(row=r, column=c).value for c in range(1, n_cols + 1)]
            for r in range(1, n_rows + 1)]

    header_row = _find_header_row(grid)
    label_col = _find_label_col(grid, header_row)
    last_row = _data_end(grid, header_row)

    col_headers = [str(grid[header_row - 1][c - 1] or f"Column {get_column_letter(c)}")
                   for c in range(1, n_cols + 1)]
    row_labels = [str(grid[r - 1][label_col - 1] or f"Row {r}")
                  for r in range(header_row + 1, last_row + 1)]

    cells: list[Cell] = []
    for r in range(1, last_row + 1):
        for c in range(1, n_cols + 1):
            cf, cv = sf.cell(row=r, column=c), sv.cell(row=r, column=c)
            raw = cv.value
            error = raw if isinstance(raw, str) and raw.startswith("#") else None
            formula = cf.value if isinstance(cf.value, str) and cf.value.startswith("=") else None
            fmt = getattr(cf, "number_format", None)
            font = getattr(cf, "font", None)

            try:
                comment = cf.comment.text.strip() if cf.comment else None
            except Exception:
                comment = None

            cells.append(Cell(
                ref=f"{get_column_letter(c)}{r}", row=r, col=c,
                value=None if error else raw,
                display=error or _format(raw, fmt),
                row_label=(str(grid[r - 1][label_col - 1]) if r > header_row
                           and grid[r - 1][label_col - 1] is not None else None),
                col_header=(col_headers[c - 1] if r != header_row else None),
                fill=_fill_of(cf),
                font_color=_font_colour_of(cf),
                bold=bool(font and font.bold),
                italic=bool(font and font.italic),
                underline=bool(font and font.underline),
                strike=bool(font and font.strike),
                bordered=_has_border(cf),
                comment=comment,
                number_format=fmt,
                formula=formula,
                error=error,
            ))

    _mark_conditional(sf, cells)

    title = _title(sf, name, grid, header_row)
    # Title first: a column header like "Growth vs Jan" names one column, not the sheet.
    value_label = _value_label(title, name, col_headers)
    model = SheetModel(
        id=f"{name}-{last_row}x{n_cols}",
        title=title,
        n_rows=last_row, n_cols=n_cols,
        header_row=header_row, label_col=label_col,
        col_headers=col_headers, row_labels=row_labels,
        value_label=value_label,
        overview="", cells=cells, attention_order=[],
    )

    # Charts are extra. One we cannot read leaves the sheet exactly as it was.
    try:
        model.charts = read_charts(formulas, values, name, cells, _format)
    except Exception:
        model.charts = []
    return model, _step_word(col_headers)


def _title(sheet, name: str, grid: list[list[Any]], header_row: int) -> str:
    """The sheet's name, unless A1 is a lone title sitting above the table."""
    try:
        first = grid[0][0]
        if header_row > 1 and isinstance(first, str) and first.strip():
            rest = [v for v in grid[0][1:] if v is not None and str(v).strip()]
            if not rest:
                return first.strip()
    except Exception:
        pass
    return name


def _mark_conditional(sheet, cells: list[Cell]) -> None:
    """Flag cells covered by a conditional formatting rule.

    We do not evaluate the rules -- that is a whole spreadsheet engine. But knowing
    the author set one here is itself information a sighted user gets for free.
    """
    try:
        covered: set[str] = set()
        for rng in sheet.conditional_formatting:
            for block in str(rng.sqref).split():
                for row in sheet[block]:
                    for cell in (row if isinstance(row, tuple) else (row,)):
                        covered.add(cell.coordinate)
        for cell in cells:
            if cell.ref in covered:
                cell.conditional = True
    except Exception:
        pass
