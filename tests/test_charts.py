"""Chart extraction: the one line chart in the chart demo, and failing safely.

The demo file was saved by Excel itself (see scripts/make_chart_demo.py), so these
tests read Excel's own chart XML. The generated workbooks below are written by
openpyxl, which stores no cached values: they prove charts are read from the cells
they cite, not from whatever copy happens to be inside the chart.
"""

import importlib.util
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "api"))

from _lib import charts, extract as extract_module  # noqa: E402
from _lib.extract import _format, extract  # noqa: E402

CHART_DEMO = ROOT / "public" / "demo" / "sales_chart_demo.xlsx"
DEMO = ROOT / "public" / "demo" / "sales_demo.xlsx"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
ITALY = [100000, 110000, 121000, 132000, 141000, 147000,
         154000, 62000, 69000, 76000, 82000, 89000]


# --- the chart demo ------------------------------------------------------------

@pytest.fixture(scope="module")
def demo():
    model, _ = extract(CHART_DEMO)
    return model


def test_the_chart_demo_has_exactly_one_line_chart(demo):
    assert len(demo.charts) == 1
    assert demo.charts[0].kind == "line"
    assert demo.charts[0].id == "chart1"


def test_title_and_axis_titles_are_read(demo):
    chart = demo.charts[0]
    assert chart.title == "Italy: Monthly Sales 2026"
    assert chart.x_title == "Month"
    assert chart.y_title == "Sales"


def test_the_anchor_says_where_the_chart_sits(demo):
    assert demo.charts[0].anchor == "B12"


def test_the_series_name_and_source_are_read(demo):
    [series] = demo.charts[0].series
    assert series.name == "Italy"
    assert series.source == "'European Sales 2026'!$B$2:$M$2"


def test_categories_values_and_source_cells_are_read(demo):
    points = demo.charts[0].series[0].points
    assert [p.category for p in points] == MONTHS
    assert [p.value for p in points] == [float(v) for v in ITALY]
    assert [p.ref for p in points] == [f"{col}2" for col in "BCDEFGHIJKLM"]


def test_each_point_matches_the_grid_cell_it_cites(demo):
    cells = {c.ref: c for c in demo.cells}
    for point in demo.charts[0].series[0].points:
        assert point.value == cells[point.ref].value


def test_every_point_reads_in_the_series_own_format(demo):
    # J2 is formatted as a percentage and K2:M2 as plain numbers in the sheet (the
    # planted formatting cues). The chart still shows them all as euros.
    displays = [p.display for p in demo.charts[0].series[0].points]
    assert displays == [f"€{v:,}" for v in ITALY]


def test_excel_cached_a_copy_of_the_values_we_can_fall_back_on():
    workbook = load_workbook(CHART_DEMO)
    values = load_workbook(CHART_DEMO, data_only=True)
    source = workbook.worksheets[0]._charts[0].series[0].val.numRef
    reader = charts._Reader(workbook, values, workbook.sheetnames[0], {}, _format)
    assert [value for _, value, _, _ in reader.cached(source)] == [float(v) for v in ITALY]


def test_the_chart_does_not_change_the_cell_cues():
    with_chart, _ = extract(CHART_DEMO)
    without, _ = extract(DEMO)
    assert [(c.ref, c.value, c.display) for c in with_chart.cells] == \
           [(c.ref, c.value, c.display) for c in without.cells]


def test_the_working_demo_has_no_charts():
    model, _ = extract(DEMO)
    assert model.charts == []


def test_upload_returns_the_chart():
    spec = importlib.util.spec_from_file_location("api.index", ROOT / "api" / "index.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    client = TestClient(module.app)
    response = client.post("/api/upload", files={
        "file": ("sales_chart_demo.xlsx", CHART_DEMO.read_bytes(), XLSX)})
    body = response.json()
    assert response.status_code == 200
    assert [c["title"] for c in body["charts"]] == ["Italy: Monthly Sales 2026"]
    assert body["attention_order"] == ["J1", "F5", "N6", "K1", "I2", "D4"]


# --- generated workbooks -----------------------------------------------------------

def _sheet(row_values=ITALY):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["Country"] + MONTHS[:len(row_values)])
    sheet.append(["Italy"] + list(row_values))
    sheet.append(["France"] + [v + 5000 if isinstance(v, int) else 1 for v in row_values])
    return workbook, sheet


def _line(sheet, row=2, anchor="B6", title="Italy"):
    chart = LineChart()
    chart.title = title
    last = sheet.max_column
    chart.add_data(Reference(sheet, min_col=1, max_col=last, min_row=row, max_row=row),
                   from_rows=True, titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=2, max_col=last, min_row=1, max_row=1))
    sheet.add_chart(chart, anchor)
    return chart


def _save(workbook, tmp_path, name="book.xlsx"):
    path = tmp_path / name
    workbook.save(path)
    return path


def test_a_chart_without_cached_values_is_read_from_its_cells(tmp_path):
    workbook, sheet = _sheet()
    _line(sheet)
    model, _ = extract(_save(workbook, tmp_path))
    [chart] = model.charts
    assert chart.anchor == "B6"
    assert [p.value for p in chart.series[0].points] == [float(v) for v in ITALY]


def test_blank_text_and_error_cells_become_gaps(tmp_path):
    workbook, sheet = _sheet([100, None, "n/a", 130, 140])
    sheet["C2"] = None
    sheet["E2"] = "#DIV/0!"
    _line(sheet)
    model, _ = extract(_save(workbook, tmp_path))
    values = [p.value for p in model.charts[0].series[0].points]
    assert values == [100.0, None, None, None, 140.0]


def test_only_line_charts_are_read(tmp_path):
    workbook, sheet = _sheet()
    bar = BarChart()
    bar.add_data(Reference(sheet, min_col=2, max_col=13, min_row=2), from_rows=True)
    sheet.add_chart(bar, "B6")
    model, _ = extract(_save(workbook, tmp_path))
    assert model.charts == []


def test_only_the_first_line_chart_is_read(tmp_path):
    workbook, sheet = _sheet()
    _line(sheet, row=2, anchor="B6", title="First")
    _line(sheet, row=3, anchor="B30", title="Second")
    model, _ = extract(_save(workbook, tmp_path))
    assert [c.title for c in model.charts] == ["First"]


def test_a_chart_pointing_at_a_missing_sheet_is_skipped_safely(tmp_path):
    workbook, sheet = _sheet()
    data = workbook.create_sheet("Data")
    data.append(MONTHS)
    data.append(ITALY)
    chart = LineChart()
    chart.add_data(Reference(data, min_col=1, max_col=12, min_row=2), from_rows=True)
    sheet.add_chart(chart, "B6")
    workbook.remove(data)     # the chart now cites a sheet that is gone, with no cache
    model, _ = extract(_save(workbook, tmp_path))
    assert model.charts == []
    assert [c.value for c in model.cells if c.row == 2][1:] == ITALY


def test_a_sheet_without_charts_has_none(tmp_path):
    workbook, _ = _sheet()
    model, _ = extract(_save(workbook, tmp_path))
    assert model.charts == []


def test_a_chart_reader_crash_leaves_the_sheet_untouched(tmp_path, monkeypatch):
    workbook, sheet = _sheet()
    _line(sheet)
    path = _save(workbook, tmp_path)
    before, _ = extract(path)

    def explode(*args, **kwargs):
        raise RuntimeError("unreadable chart")

    monkeypatch.setattr(extract_module, "read_charts", explode)
    after, _ = extract(path)
    assert after.charts == []
    assert [c.model_dump() for c in after.cells] == [c.model_dump() for c in before.cells]
