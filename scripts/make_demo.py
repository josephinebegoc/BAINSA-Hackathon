"""Build the demo workbook: 8 European countries x 12 months of revenue.

Run locally, never on Vercel:

    ./.venv/bin/python scripts/make_demo.py

That writes build/sales_demo_raw.xlsx. openpyxl cannot calculate formulas, so the
raw file has no cached values and the planted #DIV/0! does not exist yet. It has
to be recalculated by a real spreadsheet app before it is any use as a demo:

    open build/sales_demo_raw.xlsx in Excel, save it, and copy the result to
    public/demo/sales_demo.xlsx

(LibreOffice would do it headless with
 `soffice --headless --convert-to xlsx --outdir build/recalc build/sales_demo_raw.xlsx`
 but it is not installed on this machine.)

The numbers here are deliberately identical to public/fixtures/sample_sheet.json:
same seed, same planted cells. So the engine's output on this file should match the
fixture value-for-value, not merely in shape.
"""

import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "build" / "sales_demo_raw.xlsx"

TITLE = "Sales Performance 2026"
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
COL_HEADERS = ["Country"] + MONTHS + ["Growth %"]

COUNTRIES = [
    ("France", 78_000), ("Italy", 80_000), ("Germany", 100_000), ("Spain", 62_000),
    ("Netherlands", 54_000), ("Belgium", 41_000), ("Poland", 36_000), ("Sweden", 47_000),
]

SEED = 20260926          # must match the fixture
RED = "FFFF0000"


def build_revenue() -> dict[str, list[int | None]]:
    """Gentle upward drift with tight noise, then the four planted moments."""
    rng = random.Random(SEED)
    revenue: dict[str, list[int | None]] = {}
    for country, base in COUNTRIES:
        vals: list[int | None] = []
        for m in range(12):
            drift = 1.0 + 0.015 * m
            noise = 1.0 + rng.uniform(-0.035, 0.035)
            vals.append(int(round(base * drift * noise, -2)))
        revenue[country] = vals

    # 3. Trend reversal: four months of climbing, then a cliff (Germany, Jun).
    revenue["Germany"][:6] = [100_000, 110_000, 120_000, 135_000, 145_000, 80_000]
    revenue["Germany"][6:] = [92_000, 99_000, 105_000, 112_000, 118_000, 124_000]

    # 2. Anomaly: a collapse with no formatting at all (Italy, May).
    #    Feb is nudged down so May is an anomaly only, not also a trend reversal
    #    (a reversal needs three consecutive rises in front of it).
    revenue["Italy"][1] = revenue["Italy"][0] - 3_100
    revenue["Italy"][4] = 31_000

    # 4. Formula error: no November figure, so Growth % divides by an empty cell.
    revenue["Poland"][10] = None

    return revenue


def main() -> None:
    revenue = build_revenue()

    wb = Workbook()
    ws = wb.active
    ws.title = TITLE

    for c, head in enumerate(COL_HEADERS, start=1):
        ws.cell(row=1, column=c, value=head).font = Font(bold=True)

    for r, (country, _) in enumerate(COUNTRIES, start=2):
        ws.cell(row=r, column=1, value=country)

        for i, month in enumerate(MONTHS):
            value = revenue[country][i]
            if value is None:              # Poland's missing November
                continue
            cell = ws.cell(row=r, column=2 + i, value=value)
            cell.number_format = '"€"#,##0'
            # 1. Visual: one ordinary value highlighted red (Spain, March).
            if country == "Spain" and month == "Mar":
                cell.fill = PatternFill(fill_type="solid", fgColor=RED)

        growth = ws.cell(row=r, column=14, value=f"=(M{r}-L{r})/L{r}")
        growth.number_format = "0.0%"

    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 14

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)

    print(f"wrote {OUT.relative_to(ROOT)}")
    print()
    print("Planted signals, once recalculated:")
    print("  D5  Spain    Mar       red fill   -> visual")
    print("  F3  Italy    May       €31,000    -> anomaly")
    print("  G4  Germany  Jun       €80,000    -> trend reversal")
    print("  N8  Poland   Growth %  #DIV/0!    -> error")
    print()
    print("NEXT, by hand (openpyxl cannot calculate formulas):")
    print("  1. Open build/sales_demo_raw.xlsx in Excel")
    print("  2. Save it (Cmd+S), then close")
    print("  3. Copy it to public/demo/sales_demo.xlsx")
    print()
    print("Without that, the Growth % column has no cached values and the")
    print("#DIV/0! error signal simply will not exist.")


if __name__ == "__main__":
    main()
