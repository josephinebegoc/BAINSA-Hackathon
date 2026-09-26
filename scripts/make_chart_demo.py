"""Build the chart demo: the working demo sheet plus one line chart.

    python scripts/make_chart_demo.py

Reads public/demo/sales_demo.xlsx (never writes to it), adds a line chart of
Italy's monthly sales below the table, and writes build/sales_chart_demo_raw.xlsx.

Then open that file in Excel and save it as public/demo/sales_chart_demo.xlsx, so
the chart in the demo is Excel's own XML, with the values Excel caches inside it,
rather than openpyxl's. That is the kind of file a real user uploads. On Windows
with Excel installed, --excel does the open-and-save step for you.

Italy is the row with a story: seven months of growth, then a fall in August.
"""

import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "public" / "demo" / "sales_demo.xlsx"
RAW = ROOT / "build" / "sales_chart_demo_raw.xlsx"
OUT = ROOT / "public" / "demo" / "sales_chart_demo.xlsx"

ROW = 2            # Italy
FIRST_COL = 2      # Jan (B)
LAST_COL = 13      # Dec (M)
ANCHOR = "B12"     # below the table, which ends at row 9


def build() -> None:
    workbook = load_workbook(SOURCE)
    sheet = workbook.worksheets[0]

    chart = LineChart()
    chart.title = "Italy: Monthly Sales 2026"
    chart.x_axis.title = "Month"
    chart.y_axis.title = "Sales"
    # openpyxl 3.1 marks axes as deleted by default, which hides them in Excel.
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.width, chart.height = 22, 9

    data = Reference(sheet, min_col=1, max_col=LAST_COL, min_row=ROW, max_row=ROW)
    chart.add_data(data, from_rows=True, titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=FIRST_COL, max_col=LAST_COL,
                                   min_row=1, max_row=1))
    sheet.add_chart(chart, ANCHOR)

    RAW.parent.mkdir(exist_ok=True)
    workbook.save(RAW)
    print(f"wrote {RAW.relative_to(ROOT)}")


def resave_with_excel() -> None:
    """Open the raw file in Excel and save it as the demo (Windows + Excel only)."""
    script = f"""
$excel = New-Object -ComObject Excel.Application
$excel.Visible = $false
$excel.DisplayAlerts = $false
try {{
  $book = $excel.Workbooks.Open('{RAW}')
  $book.SaveAs('{OUT}', 51)
  $book.Close($false)
}} finally {{
  $excel.Quit()
  [void][Runtime.InteropServices.Marshal]::ReleaseComObject($excel)
}}
"""
    subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
                   check=True)
    print(f"Excel saved {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    build()
    if "--excel" in sys.argv:
        resave_with_excel()
    elif "--copy" in sys.argv:
        # No Excel available: use openpyxl's file as it is. It still works, because
        # chart values are read from the cells, but it is not what Excel writes.
        shutil.copyfile(RAW, OUT)
        print(f"copied to {OUT.relative_to(ROOT)} (not re-saved by Excel)")
    else:
        print(f"now open it in Excel and save it as {OUT.relative_to(ROOT)}")
